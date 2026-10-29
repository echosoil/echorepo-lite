#!/usr/bin/env python3
"""Safely delete Firestore sample DOCUMENTS by QR codes in CSV/XLSX/TXT.

FireStore layout (as used by batch_upload_or_update_samples.py):
    users/<uid>/samples/<sampleDocId>
    sample["data"][*] = {"type": "QR", "info": {"qrCode": "ABCD-1234"}, ...}

DRY RUN is the default.  --apply requires --confirm-project and a separate,
explicit typed confirmation for EVERY document.  Full document backups are
written in Firestore protobuf format BEFORE any deletion.  No Firebase Storage
objects, Auth users, or MinIO objects are deleted.

Examples:
  python3 tools/delete_firestore_samples_by_qr.py codes.txt
  python3 tools/delete_firestore_samples_by_qr.py codes.xlsx --sheet qr_codes
  python3 tools/delete_firestore_samples_by_qr.py codes.csv --apply --confirm-project echosoil
  python3 tools/delete_firestore_samples_by_qr.py --restore-from data/maintenance/qr_deletions/backup_....jsonl
  python3 tools/delete_firestore_samples_by_qr.py --restore-from ... --apply --confirm-project echosoil
"""
from __future__ import annotations

import argparse
import base64
import csv
import hashlib
import json
import os
import re
import sys
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4

ROOT = Path(__file__).resolve().parent.parent
QR_COLUMNS = ("QR Code", "QR_qrCode", "qrCode", "qr_code", "qr", "QR", "code")
QR_PATTERN = re.compile(r"^([A-Z]{4})-?([0-9]{4})$")


def norm_qr(value):
    """Conservative normalization: whitespace, case and hyphen variants only."""
    if value is None:
        return None
    s = str(value).strip().upper()
    for dash in ("–", "—", "−", "‐", "‑", "‒"):
        s = s.replace(dash, "-")
    s = re.sub(r"\s+", "", s)
    hit = QR_PATTERN.fullmatch(s)
    return f"{hit[1]}-{hit[2]}" if hit else None


def read_codes(path: Path, sheet: str | None = None):
    """Return unique QR codes in file order, and rejected nonblank entries."""
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"Input not found: {path}")
    suffix = path.suffix.lower()
    values = []
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet:
                if sheet not in wb.sheetnames:
                    raise ValueError(f"Sheet {sheet!r} not found: {wb.sheetnames}")
                ws = wb[sheet]
            else:
                ws = next((w for w in wb.worksheets if w.title.lower() != "ignoreme"), None)
                if ws is None:
                    raise ValueError("No usable workbook sheet")
            rows = ws.iter_rows(values_only=True)
            header = next(rows, None)
            if not header:
                return [], []
            col = find_qr_column(header)
            if col is None:
                if len(header) == 1 and norm_qr(header[0]):
                    col = 0
                    values.append((1, header[0]))
                elif len(header) == 1:
                    raise ValueError(f"Unrecognized QR header {header[0]!r}; use one of {QR_COLUMNS}")
                else:
                    raise ValueError(f"No QR column found: {header}")
            for i, row in enumerate(rows, start=2):
                values.append((i, row[col] if col < len(row) else None))
        finally:
            wb.close()
    elif suffix in {".csv", ".txt"}:
        with path.open(encoding="utf-8-sig", newline="") as fh:
            if suffix == ".txt":
                values = [(i, line) for i, line in enumerate(fh, 1)]
            else:
                sample = fh.read(8192)
                fh.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
                    delimiter = dialect.delimiter
                except csv.Error:
                    delimiter = ","
                reader = csv.reader(fh, delimiter=delimiter)
                header = next(reader, None)
                if not header:
                    return [], []
                col = find_qr_column(header)
                if col is None:
                    if len(header) == 1 and norm_qr(header[0]):
                        col = 0
                        values.append((1, header[0]))
                    else:
                        raise ValueError(f"No QR column found; expected one of {QR_COLUMNS}; got {header}")
                for i, row in enumerate(reader, start=2):
                    values.append((i, row[col] if col < len(row) else None))
    else:
        raise ValueError("Use CSV, XLSX, XLSM, or plain TXT (one QR per line)")

    codes = []
    bad = []
    seen = set()
    for line, value in values:
        raw = str(value).strip() if value is not None else ""
        if not raw or raw.startswith("#"):
            continue
        qr = norm_qr(raw)
        if qr is None:
            bad.append((line, raw))
        elif qr not in seen:
            seen.add(qr)
            codes.append(qr)
    return codes, bad


def find_qr_column(headers):
    names = [str(item or "").strip().casefold() for item in headers]
    for name in QR_COLUMNS:
        if name.casefold() in names:
            return names.index(name.casefold())
    return None


def sample_qrs(doc_data):
    """Read QR values ONLY from the QR step (not from unrelated text)."""
    matches = set()
    steps = doc_data.get("data") if isinstance(doc_data, dict) else None
    if not isinstance(steps, list):
        return matches
    for step in steps:
        if not isinstance(step, dict) or step.get("type") != "QR":
            continue
        info = step.get("info")
        if isinstance(info, dict):
            qr = norm_qr(info.get("qrCode"))
            if qr:
                matches.add(qr)
    return matches


def allowed_path(path):
    pieces = str(path).split("/")
    return len(pieces) == 4 and pieces[0] == "users" and pieces[2] == "samples" and bool(pieces[1]) and bool(pieces[3])


def secure_write(path: Path, records):
    path.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
    if os.name == "posix":
        os.chmod(path.parent, 0o700)
    fd = os.open(path, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="") as f:
            for record in records:
                f.write(json.dumps(record, ensure_ascii=False, allow_nan=False, sort_keys=True) + "\n")
            f.flush()
            os.fsync(f.fileno())
    except Exception:
        path.unlink(missing_ok=True)
        raise


def serialize_snapshot(snap):
    """Lossless backup including Firestore types (timestamps, GeoPoints, bytes).

    Encodes the native Firestore Document protobuf to base64; if the encoder
    cannot handle a value, abort BEFORE deleting anything rather than writing
    an incomplete JSON backup.
    """
    from google.cloud.firestore_v1 import _helpers
    from google.cloud.firestore_v1.types import Document
    pb = _helpers.document_snapshot_to_protobuf(snap)
    if pb is None:
        raise ValueError(f"Document disappeared during backup: {snap.reference.path}")
    raw = Document.pb(pb).SerializeToString()
    return base64.b64encode(raw).decode("ascii"), hashlib.sha256(raw).hexdigest()


def deserialize_snapshot(blob, expected_sha, client):
    from google.cloud.firestore_v1 import _helpers
    from google.cloud.firestore_v1.types import Document
    raw = base64.b64decode(blob, validate=True)
    if hashlib.sha256(raw).hexdigest() != expected_sha:
        raise ValueError("Backup entry checksum mismatch")
    pb = Document.deserialize(raw)
    return _helpers.decode_dict(pb.fields, client)


def firebase_connection(service_account: str | None):
    import firebase_admin
    from firebase_admin import credentials, firestore
    from dotenv import load_dotenv

    load_dotenv(ROOT / ".env")
    paths = []
    if service_account:
        paths.append(Path(service_account).expanduser())
    elif os.getenv("GOOGLE_APPLICATION_CREDENTIALS"):
        paths.append(Path(os.environ["GOOGLE_APPLICATION_CREDENTIALS"]).expanduser())
    else:
        # Mirror pull_and_enrich_samples.py, with a fallback for host execution.
        paths.extend([Path("/opt/echorepo/keys/firebase-sa.json"),
                      ROOT / "opt/echorepo/keys/firebase-sa.json",
                      ROOT / "../../../keys/firebase-sa.json"])
    resolved_paths = []
    for candidate in paths:
        if candidate.is_absolute():
            resolved_paths.append(candidate)
            # Same host-side fallback as pull_and_enrich_samples.py when
            # an absolute /opt/... path exists only inside Docker.
            resolved_paths.append(ROOT / candidate.relative_to("/"))
        else:
            resolved_paths.append(ROOT / candidate)
    credential_path = next((p for p in resolved_paths if p.is_file()), None)
    if credential_path is None:
        raise FileNotFoundError("Service account not found. Set GOOGLE_APPLICATION_CREDENTIALS in .env or pass --service-account")
    sa = json.loads(credential_path.read_text(encoding="utf-8"))
    project_sa = str(sa.get("project_id") or "")
    project_env = str(os.getenv("FIREBASE_PROJECT_ID") or "")
    if project_sa and project_env and project_sa != project_env:
        raise ValueError("FIREBASE_PROJECT_ID and service-account project_id differ")
    project = project_env or project_sa
    if not project:
        raise ValueError("Firebase project cannot be determined")
    try:
        app = firebase_admin.get_app()
        if app.project_id != project:
            raise RuntimeError("Existing Firebase app belongs to a different project")
    except ValueError:
        app = firebase_admin.initialize_app(credentials.Certificate(str(credential_path)), {"projectId": project})
    return firestore.client(app), project


def discover(db, codes, only_user=None):
    wanted = set(codes)
    hits = defaultdict(list)
    scanned = 0
    print("Scanning Firestore collection_group('samples')...", flush=True)
    for snap in db.collection_group("samples").stream(timeout=300):
        scanned += 1
        path = snap.reference.path
        if not allowed_path(path):
            continue
        if only_user and path.split("/")[1] != only_user:
            continue
        common = sample_qrs(snap.to_dict() or {}) & wanted
        if common:
            # Store ONE entry per doc/QR; never silently guess owners via lookup.
            for code in common:
                hits[code].append(snap)
        if scanned % 1000 == 0:
            print(f"  scanned {scanned} documents...", flush=True)
    print(f"Scanned {scanned} documents; found {sum(map(len,hits.values()))} matching documents")
    return hits, scanned


def ask_for_confirmation(action, qr, path):
    print(f"\n  QR:       {qr}")
    print(f"  Document: {path}")
    print(f"  To {action.lower()} THIS document, type: {action} {qr}")
    try:
        answer = input("  Confirmation (Enter to skip): ").strip()
    except (KeyboardInterrupt, EOFError):
        print("\nCancelled; remaining documents unchanged.")
        return False
    return answer == f"{action} {qr}"


def prepare_records(project, codes, found):
    """Prepare a complete backup and human-readable plan before ANY mutation."""
    backup, plan = [], []
    for qr in codes:
        matches = sorted(found.get(qr, []), key=lambda s: s.reference.path)
        if not matches:
            plan.append({"qr": qr, "status": "NOT_FOUND", "document_path": ""})
            continue
        for snap in matches:
            path = snap.reference.path
            if not snap.exists:
                raise RuntimeError(f"Document vanished while preparing plan: {path}")
            encoded, digest = serialize_snapshot(snap)
            backup.append({
                "schema_version": 1,
                "firebase_project": project,
                "qr": qr,
                "document_path": path,
                "scanned_update_time": str(snap.update_time),
                "firestore_document_proto_b64": encoded,
                "sha256": digest,
            })
            plan.append({"qr": qr, "status": "MULTIPLE_MATCHES" if len(matches) > 1 else "FOUND",
                         "document_path": path})
    return backup, plan


def apply_deletions(db, codes, found, allow_duplicates=False):
    # Let Firestore compare the scanned document update_time atomically.
    # Comparing Firestore timestamp wrappers in Python can produce false
    # mismatches (e.g. different nanosecond precision/representations).
    from google.api_core.exceptions import FailedPrecondition, NotFound

    removed = skipped = changed = failed = 0
    for qr in codes:
        snaps = sorted(found.get(qr, []), key=lambda s: s.reference.path)
        if not snaps:
            print(f"NOT FOUND {qr}")
            continue
        if len(snaps) > 1 and not allow_duplicates:
            print(f"SKIP {qr}: {len(snaps)} matching documents; use --allow-duplicate-qrs after checking the plan")
            skipped += len(snaps)
            continue
        for scanned in snaps:
            path = scanned.reference.path
            print(f"\nCandidate: {path} (QR {qr})")
            if not ask_for_confirmation("DELETE", qr, path):
                print("  SKIPPED (not confirmed)")
                skipped += 1
                continue
            try:
                ref = db.document(path)
                latest = ref.get()
                if not latest.exists:
                    print("  SKIP: document was deleted after the scan")
                    changed += 1
                    continue
                if qr not in sample_qrs(latest.to_dict() or {}):
                    print("  SKIP: QR field changed after the scan")
                    changed += 1
                    continue
                # The scanned timestamp is the *original* version being backed up.
                # Firestore itself enforces this condition atomically at deletion;
                # never compare latest.update_time and scanned.update_time in Python.
                ref.delete(option=db.write_option(last_update_time=scanned.update_time))
                removed += 1
                print(f"  DELETED {path}")
            except (FailedPrecondition, NotFound) as exc:
                # A concurrent change/deletion was rejected by Firestore.
                # No Firestore document was changed by this script.
                changed += 1
                print(f"  SKIP: Firestore rejected a changed/deleted document ({type(exc).__name__})")
            except Exception as exc:
                failed += 1
                print(f"  ERROR {path}: {type(exc).__name__}: {exc}", file=sys.stderr)
    print(f"\nResult: deleted={removed} skipped={skipped} changed={changed} failed={failed}")
    return 1 if failed or changed else 0


def restore_deletions(db, project, backup_path: Path, apply: bool):
    """Restore missing documents, never replace ones that already exist."""
    records = []
    with backup_path.open(encoding="utf-8") as f:
        for line in f:
            if line.strip():
                entry = json.loads(line)
                if entry.get("firebase_project") != project:
                    raise ValueError("Backup belongs to another Firebase project")
                if entry.get("schema_version") != 1 or not allowed_path(entry.get("document_path", "")):
                    raise ValueError("Unrecognized backup record")
                records.append(entry)
    if not records:
        print("Backup is empty")
        return 0
    # Verify ALL backups before any restore begins.
    decoded = [(rec, deserialize_snapshot(rec["firestore_document_proto_b64"], rec["sha256"], db))
               for rec in records]
    print(f"Restore candidates: {len(decoded)} documents; project={project}")
    for rec, data in decoded:
        path = rec["document_path"]
        print(f"RESTORE CANDIDATE {rec['qr']}: {path}")
        if not apply:
            continue
        if not ask_for_confirmation("RESTORE", rec["qr"], path):
            print("  SKIPPED")
            continue
        ref = db.document(path)
        if ref.get().exists:
            print("  SKIP: document already exists")
            continue
        # create() is exclusive; concurrent recreation fails, not overwritten.
        ref.create(data)
        print(f"  RESTORED {path}")
    if not apply:
        print("DRY RUN: no Firestore documents restored")
    return 0


def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("input_file", type=Path, nargs="?", help="CSV/XLSX/TXT with QR codes")
    p.add_argument("--sheet", help="Excel sheet (default: first sheet except ignoreme)")
    p.add_argument("--user-id", help="Restrict to samples owned by one Firebase UID")
    p.add_argument("--service-account", help="Firebase service-account JSON (otherwise .env)")
    p.add_argument("--apply", action="store_true", help="Actually delete documents; default is dry run")
    p.add_argument("--confirm-project", help="Must match project ID for --apply")
    p.add_argument("--allow-duplicate-qrs", action="store_true", help="Allow deletion of ambiguous QRs, still one confirmation per document")
    p.add_argument("--report-dir", type=Path, default=ROOT / "data/maintenance/qr_deletions",
                   help="Directory for private full-document backup and plan")
    p.add_argument("--restore-from", type=Path, help="Restore missing documents from a previous backup (dry run by default)")
    args = p.parse_args(argv)
    if bool(args.input_file) == bool(args.restore_from):
        p.error("Provide an input file OR --restore-from, not both")
    if args.apply and not args.confirm_project:
        p.error("--apply requires --confirm-project PROJECT_ID")
    if args.restore_from and (args.user_id or args.sheet or args.allow_duplicate_qrs):
        p.error("--restore-from cannot be combined with --user-id/--sheet/--allow-duplicate-qrs")

    codes = bad = None
    if args.input_file:
        codes, bad = read_codes(args.input_file, sheet=args.sheet)
        print(f"Input: {len(codes)} distinct valid QR codes, {len(bad)} invalid nonempty entries")
        for line, text in bad[:25]:
            print(f"  INVALID row {line}: {text!r}")
        if not codes:
            raise ValueError("No valid QR codes found; nothing to do")
        if bad:
            print("WARNING: invalid rows will not be searched/deleted")

    db, project = firebase_connection(args.service_account)
    print(f"Firebase project: {project}")
    if args.apply and args.confirm_project != project:
        p.error(f"Project mismatch: --confirm-project must equal {project!r}")

    if args.restore_from:
        return restore_deletions(db, project, args.restore_from, args.apply)

    matches, _ = discover(db, codes, args.user_id)
    backup, plan = prepare_records(project, codes, matches)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ") + "_" + uuid4().hex[:8]
    outdir = args.report_dir if args.report_dir.is_absolute() else ROOT / args.report_dir
    backup_path = outdir / f"backup_{stamp}.jsonl"
    plan_path = outdir / f"plan_{stamp}.jsonl"
    secure_write(backup_path, backup)
    secure_write(plan_path, plan)
    print(f"Private FULL document backup: {backup_path}")
    print(f"Private plan: {plan_path}")
    print("Requested QRs:")
    for qr in codes:
        paths = [x.reference.path for x in matches.get(qr, [])]
        state = "NOT FOUND" if not paths else f"{len(paths)} document(s)"
        print(f"  {qr}: {state}" + (" -> " + ", ".join(paths) if paths else ""))
    if not args.apply:
        print("DRY RUN: no deletion. To delete, use --apply --confirm-project PROJECT_ID")
        return 0
    return apply_deletions(db, codes, matches, args.allow_duplicate_qrs)


if __name__ == "__main__":
    try:
        sys.exit(main())
    except (KeyboardInterrupt, EOFError):
        print("\nCancelled. Remaining documents unchanged.", file=sys.stderr)
        sys.exit(130)
    except Exception as exc:
        print(f"ERROR: {type(exc).__name__}: {exc}", file=sys.stderr)
        sys.exit(1)
